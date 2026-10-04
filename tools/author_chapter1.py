# -*- coding: utf-8 -*-
"""第一章「棺女」五夜章节落表脚本（2026-10-04）。
读取 config/*.xlsx，克隆现有行并改写为 5 夜章节配置：
- battle_config.xlsx: 新增 ch1_night1 / ch1_night3 / ch1_night4 三场战斗（敌方/机关/卡池行）
- stage_config.xlsx: 新增 night3_procession / night4_mirror 舞台 + 线索，改 corridor 的 next 链；重写 night_patrol 结案 nodes
- dialogue.xlsx: 新增 corridor_act / night3_procession / night4_mirror / honden_act 四段探索对话
- data/commissions.json: 5 条委托
幂等：重复运行前先删除本脚本写入的行（按首列标记识别）。
"""
import json
import openpyxl

ROOT = __file__.rsplit('\\tools\\', 1)[0] if '\\tools\\' in __file__ else '.'


def hdrmap(ws):
    hdr = [c.value for c in ws[1]]
    m = {}
    for i, h in enumerate(hdr):
        if h is None:
            continue
        key = str(h).split('(')[0].strip()
        m.setdefault(key, i)
    return hdr, m


def col(m, *names):
    for n in names:
        if n in m:
            return m[n]
    raise KeyError(names)


# ---------------- battle_config.xlsx ----------------
def author_battle():
    wb = openpyxl.load_workbook(f'{ROOT}/config/battle_config.xlsx')
    ws = wb['battle']
    hdr, m = hdrmap(ws)
    rows = list(ws.iter_rows(values_only=True))
    demo = next(r for r in rows if r and r[0] == 'demo_corridor')
    # 幂等清理（按首列 battle 标记删除旧行）
    for name in ['battle', 'enemy', 'slot', 'card']:
        s = wb[name]
        drop = [i for i, r in enumerate(list(s.iter_rows(values_only=True))[1:], start=2)
                if r and r[0] in ('ch1_night1', 'ch1_night3', 'ch1_night4')]
        for i in reversed(drop):
            s.delete_rows(i, 1)

    ws = wb['battle']
    hdr, m = hdrmap(ws)
    rows = list(ws.iter_rows(values_only=True))
    demo = [list(r) for r in rows if r and r[0] == 'demo_corridor'][0]
    honden = [list(r) for r in rows if r and r[0] == 'honden_boss'][0]

    def setc(row, m, **kw):
        for k, v in kw.items():
            row[m[k]] = v

    spirit = lambda start: json.dumps({"start": start, "per_turn": 1, "phase2_from_turn": 3,
        "phase2_bonus": 2, "tier_attack_mult": {"weak": 0.7, "normal": 1.0, "empowered": 1.3, "berserk": 1.5},
        "weak_below": 3, "empowered_from": 6, "berserk_from": 8}, ensure_ascii=False)
    b1 = demo[:]
    setc(b1, m, **{'file': 'ch1_night1', 'id': 'ch1_night1',
                   'name': '第一夜 · 参道灯笼影（教学战）',
                   'comment': '第一章棺女·夜一：仅扑击/尖啸，认识灵墨与出牌',
                   'ap_per_turn': '5', 'spirit': spirit(3),
                   'demo_note': '教学战：纸人形只会扑击与尖啸。引导：攻击牌打立绘 / 引燃场景槽 / 鸣钟震慑。'})
    b3 = demo[:]
    setc(b3, m, **{'file': 'ch1_night3', 'id': 'ch1_night3',
                   'name': '第三夜 · 回廊深处·抬棺',
                   'comment': '第一章棺女·夜三：抬棺队列两波，封印教学',
                   'ap_per_turn': '6', 'spirit': spirit(5),
                   'demo_note': '两波压制：抬棺纸人会猛扑（重击）。灵气≥6 且阵眼≥2 时贴符可封印斩杀。'})
    b4 = honden[:]
    setc(b4, m, **{'file': 'ch1_night4', 'id': 'ch1_night4',
                   'name': '第四夜 · 本殿黄昏·棺中镜',
                   'comment': '第一章棺女·夜四：残响群 + 执棺精英（护体/执念）',
                   'ap_per_turn': '6', 'spirit': spirit(6),
                   'demo_note': '精英带棺气护体（需先破阵）与执念槽（对应槽位被激活时暴怒）。'})
    for b in (b1, b3, b4):
        ws.append(b)

    # enemy 行：克隆 demo/honden 的纸人形行再改
    ws = wb['enemy']
    hdr, m = hdrmap(ws)
    rows = [list(r) for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0]]
    demo_e0 = [list(r) for r in rows if r[0] == 'demo_corridor' and r[1] == 0][0]
    demo_e1 = [list(r) for r in rows if r[0] == 'demo_corridor' and r[1] == 1][0]
    honden_e2 = [list(r) for r in rows if r[0] == 'honden_boss' and r[1] == 2][0]

    def enemy(base, battle, idx, name, hp, atk, sprite, height, pattern, intent, pounce, howl, count, spacing, guard='', obsession_slot='', obsession_text=''):
        r = base[:]
        setc(r, m, **{'battle': battle, 'idx': idx, 'name': name,
                      'sprite': f'res://assets/chars/{sprite}', 'hp': str(hp), 'attack': str(atk),
                      'height_px': str(height), 'seal_guard': guard,
                      'obsession_slot': obsession_slot, 'obsession_text': obsession_text,
                      'pattern': json.dumps(pattern, ensure_ascii=False),
                      'intent_text': json.dumps(intent, ensure_ascii=False),
                      'pounce_mult': str(pounce) if pounce else '', 'howl_spirit': str(howl),
                      'count': str(count) if count else '', 'spacing_px': str(spacing) if spacing else ''})
        return r

    enemies = [
        enemy(demo_e0, 'ch1_night1', 0, '纸人形', 18, 4, 'enemy_paper_doll.png', 500,
              ['attack', 'attack', 'howl'], {'attack': '扑击', 'howl': '尖啸 —— 灵气上升！'}, '', 1, 2, 260),
        enemy(demo_e0, 'ch1_night3', 0, '纸人形·抬棺', 30, 7, 'enemy_paper_doll.png', 540,
              ['attack', 'pounce', 'attack', 'howl'],
              {'attack': '扑击', 'pounce': '猛扑（重击）', 'howl': '尖啸 —— 灵气上升！'}, 1.8, 2, 2, 300),
        enemy(demo_e1, 'ch1_night3', 1, '纸人形·残响', 22, 6, 'enemy_paper_doll_echo.png', 440,
              ['pounce', 'attack', 'attack'], {'attack': '扑击', 'pounce': '猛扑（重击）'}, 1.6, 2, '', ''),
        enemy(demo_e1, 'ch1_night4', 0, '纸人形·残响', 24, 6, 'enemy_paper_doll_echo.png', 440,
              ['pounce', 'attack', 'howl', 'attack'],
              {'attack': '扑击', 'pounce': '猛扑（重击）', 'howl': '尖啸 —— 灵气上升！'}, 1.6, 2, 2, 260),
        enemy(demo_e0, 'ch1_night4', 1, '纸人形·执棺', 36, 8, 'enemy_paper_doll.png', 560,
              ['attack', 'attack', 'pounce', 'howl'],
              {'attack': '扑击', 'pounce': '猛扑（重击）', 'howl': '尖啸 —— 灵气上升！'}, 1.8, 2, '', '',
              '18', 'slot_shoji', '执念：守棺 —— 仪式未成，棺不可开'),
    ]
    for e in enemies:
        ws.append(e)

    # slot 行：克隆 demo 5 槽给三场新战斗
    ws = wb['slot']
    hdr, m = hdrmap(ws)
    demo_slots = [list(r) for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0] == 'demo_corridor']
    for battle in ('ch1_night1', 'ch1_night3', 'ch1_night4'):
        for s in demo_slots:
            r = s[:]
            r[m['battle']] = battle
            ws.append(r)

    # card 行：夜一 8 张教学卡；夜三 = demo 全套；夜四 = honden 全套
    ws = wb['card']
    hdr, m = hdrmap(ws)
    all_cards = [list(r) for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0]]
    night1_ids = ['slash', 'ignite', 'bell', 'seal', 'seal2', 'guard', 'amulet', 'expose']
    for battle, src, ids in (
        ('ch1_night1', 'demo_corridor', night1_ids),
        ('ch1_night3', 'demo_corridor', None),
        ('ch1_night4', 'honden_boss', None),
    ):
        for c in all_cards:
            if c[0] != src:
                continue
            if ids and c[col(m, 'id')] not in ids:
                continue
            r = c[:]
            r[m['battle']] = battle
            ws.append(r)

    wb.save(f'{ROOT}/config/battle_config.xlsx')
    print('battle_config.xlsx OK')


# ---------------- stage_config.xlsx ----------------
def node(speaker, name, color, portrait, side, text, next='', choices=None):
    n = {'speaker': speaker, 'name': name, 'color': color, 'portrait': portrait, 'side': side,
         'text': text, 'next': next}
    if choices is not None:
        n['choices'] = choices
    return n

RINNE = ('rinne', '凛音', 'vermilion_500', 'res://assets/chars/portraits/rinne_half.png', 'left')
HAKUYO = ('hakuyo', '薄荷', 'fuji_500', 'res://assets/chars/portraits/hakuyo_half.png', 'right')
SAYO = ('sayo', '？？？', 'paper_300', '', 'right')
GUANSHOU = ('guanshou', '棺守', 'ink_700', '', 'right')


def author_stage():
    wb = openpyxl.load_workbook(f'{ROOT}/config/stage_config.xlsx')

    # ---- stage 表 ----
    ws = wb['stage']
    hdr, m = hdrmap(ws)
    rows = [list(r) for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0]]
    # 幂等
    drop = [i for i, r in enumerate(list(ws.iter_rows(values_only=True))[1:], start=2)
            if r and r[0] in ('night3_procession', 'night4_mirror')]
    for i in reversed(drop):
        ws.delete_rows(i, 1)
    corridor = [r for r in rows if r[0] == 'corridor_act'][0]
    honden = [r for r in rows if r[0] == 'honden_act'][0]

    def setc(row, **kw):
        for k, v in kw.items():
            row[m[k]] = v

    # corridor 的 next 改为第三夜
    for i, r in enumerate(list(ws.iter_rows(values_only=True))[1:], start=2):
        if r and r[0] == 'corridor_act':
            ws.cell(row=i, column=m['next'] + 1,
                    value='res://scenes/v3/stage_night3.tscn')
    n3 = corridor[:]
    setc(n3, **{'file': 'night3_procession', 'id': 'night3_procession',
                'name': '回廊深处（第三夜·纸人抬棺）',
                'comment': '第三夜探索：corridor_walk 深夜层，雾更重',
                'spawn_x': '560', 'spirit': '4',
                'next': 'res://scenes/v3/stage_night4.tscn'})
    n4 = honden[:]
    setc(n4, **{'file': 'night4_mirror', 'id': 'night4_mirror',
                'name': '本殿黄昏（第四夜·棺中镜）',
                'comment': '第四夜探索：honden_dusk 棺停御神木下',
                'spawn_x': '520', 'spirit': '5',
                'next': 'res://scenes/v3/stage_honden.tscn'})
    ws.append(n3)
    ws.append(n4)

    # ---- clue 表 ----
    ws = wb['clue']
    hdr, m = hdrmap(ws)
    rows = [list(r) for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0]]
    drop = [i for i, r in enumerate(list(ws.iter_rows(values_only=True))[1:], start=2)
            if r and r[0] in ('night3_procession', 'night4_mirror')]
    for i in reversed(drop):
        ws.delete_rows(i, 1)
    corridor_clue = [r for r in rows if r[0] == 'corridor_act'][0]

    def clue(stage, cid, ctype, cname, x, hint, resolve, battle, file_comment):
        r = corridor_clue[:]
        setc(r, **{'stage': stage, 'id': cid, 'type': ctype, 'name': cname,
                   'x': str(x), 'hint': hint, 'resolve_text': resolve,
                   'battle': battle, 'file_comment': file_comment})
        return r

    ws.append(clue('night3_procession', 'coffin_procession', 'visual', '抬棺队列', 700,
                   '纸人形列队而过，抬着一口空棺。棺盖缝隙透出微光——有节奏地，亮三下，熄三下。'
                   '薄荷拉住你的袖子：「它们走了七十年的同一条路，一步都没有偏过。」',
                   '空棺轻得像一页纸，又沉得压弯了抬棺的纸肩。第三声叩响落下，队列齐刷刷转向了你——'
                   '你没有让路。它们也没有。这座神社在等一个先开口的人。',
                   'res://scenes/v3/battle_night3.tscn',
                   '第三夜画面线索（第一章棺女）：夜行抬棺——棺中三叩，是催促，也是路标。'))
    ws.append(clue('night4_mirror', 'mirror_in_coffin', 'reflection', '棺中之镜', 420,
                   '棺底嵌着一面铜镜。镜中没有你的倒影——只有一个穿嫁衣的背影。'
                   '薄荷听了一会儿，轻声说：「她问，山下顾家……是不是派人来接她了。」',
                   '镜面冰凉，映出的却不是死相——是七十年前婚礼前夜，那个还在等灯的姑娘。'
                   '她的唇形很慢，像隔着很深的水：「送我走。」',
                   'res://scenes/v3/battle_night4.tscn',
                   '第四夜画面线索（第一章棺女）：棺中镜——照见小夜生前模样，顾家旧约浮出水面。'))

    # ---- case 表：重写 night_patrol 结案 nodes（五夜完整真相 + 双分支） ----
    ws = wb['case']
    hdr, m = hdrmap(ws)
    nodes = {
        't1': node(*RINNE,
                   '五夜查明白了。水钵里的灯笼、无风的钟、夜行的棺、棺中的镜——都是同一件事：'
                   '七十年前，一场没送成的葬。', 't2'),
        't2': node(*RINNE,
                   '钟鸣不是警报，是送行钟。纸人形列队不是在作乱，是在等一个替他们把仪式办完的人。'
                   '第三次叩门，是它最后的耐心。', 't3'),
        't3': node(*RINNE, '现在，轮到我了。', '', [
            {'text': '（送行）念引路词，送它远行', 'next': 'e_send'},
            {'text': '（镇守）立朱印，镇其归处', 'next': 'e_seal'}]),
        'e_send': node(*SAYO, '……谢谢你。七十年了，终于有人听懂钟声里的意思。', 's1'),
        's1': node(*SAYO, '我叫小夜。山下村子，顾家的媳妇——虽然还没来得及过门。'
                          '（她低头看自己的手）外头……仗打完了吗？', 's2'),
        's2': node(*RINNE, '打完了。七十年前就打完了。村子的人迁去了南边，开枝散叶，活得很好。'
                          '顾家的后人里，有一个做了退魔师——她让我替她还一盏灯。', 's3'),
        's3': node(*SAYO, '那就好，那就好……（她转向棺守的纸躯，端端正正福了一礼）'
                         '祭主爷爷，更次你替我数了三万次，辛苦你了。', 's4'),
        's4': node(*RINNE, '（引路词）魂兮归来，莫恋此乡。灯前有路，路尽有光。'
                          '顾氏小夜——拜堂，起灵，上路！', 's5'),
        's5': node(*SAYO, '（她终于穿上那身嫁衣的影子，就着灯火，端端正正拜了下去。'
                         '一礼既成，身影随灯而起，淡如晨雾）替我向山下问好。'
                         '就说——小夜先走一步，回有太阳的山坡去了。', 's6'),
        's6': node(*RINNE, '（棺守的纸躯委顿于地，像卸下了肩上的棺。逢魔之刻的雾，第一次在天亮前散尽了。'
                          '薄荷在哭，哭得很小声，怕吵到谁似的。）……结案。'
                          '退魔寮乙字科凛音：逢魔神社夜巡，五夜，送行者一，安魂。', ''),
        'e_seal': node(*RINNE, '（立朱印）以逢魔之名，镇其归处。自此异象不起，百夜无更。', 'z1'),
        'z1': node(*GUANSHOU, '镇……？你要……镇她？她等了七十年……'
                            '你们山下人，七十年前弃她而去，如今还要把她钉在这口棺里——？！', 'z2'),
        'z2': node(*RINNE, '（朱印落下，钟鸣、纸响、棺纹一齐寂然。只有棺底极轻极轻地，叩了一声。'
                          '像还有话，被封在了里面。）薄荷，记下来：本案未结，改「续监」。'
                          '……我总有一天，会回来把这句话听完。', 'z3'),
        'z3': node(*SAYO, '……没关系。已经，习惯了。', ''),
    }
    for i, r in enumerate(list(ws.iter_rows(values_only=True))[1:], start=2):
        if r and r[0] == 'night_patrol':
            ws.cell(row=i, column=m['nodes'] + 1,
                    value=json.dumps(nodes, ensure_ascii=False))
            ws.cell(row=i, column=m['name'] + 1, value='逢魔神社 · 夜巡（第一章「棺女」结案）')
            ws.cell(row=i, column=m['comment'] + 1,
                    value='第一章五夜真相：棺女小夜——七十年前没送成的葬；双分支结局（送行/镇守），镇守线留第二章伏笔')

    wb.save(f'{ROOT}/config/stage_config.xlsx')
    print('stage_config.xlsx OK')


# ---------------- dialogue.xlsx ----------------
def author_dialogue():
    wb = openpyxl.load_workbook(f'{ROOT}/config/dialogue.xlsx')
    ws = wb['nodes']
    # 幂等：删除本脚本写入的 stage 段
    drop = [i for i, r in enumerate(list(ws.iter_rows(values_only=True))[1:], start=2)
            if r and ((r[col({'stage': 1}, 'stage')] in ('corridor_act', 'night3_procession', 'night4_mirror', 'honden_act'))
                or (r[col({'stage': 1}, 'stage')] == 'test_approach' and r[0] in ('r1', 'r2', 'r3')))]
    for i in reversed(drop):
        ws.delete_rows(i, 1)

    def n(nid, stage, trigger, arg, who, text, next=''):
        return [nid, stage, trigger, arg, *who, text, next]

    rows = [
        # ---- 夜一参道：凛音的私缘（石阶 hotspot，独立于教学链） ----
        n('r1', 'test_approach', 'hotspot', '300', RINNE,
          '这条石阶……祖母讲过。她说小时候从山上下来，就是踩着这些台阶，一手牵着她娘，一手拎着那盏没送出去的灯。', 'r2'),
        n('r2', 'test_approach', '', '', HAKUYO,
          '没送出去的灯？……委托单上写的是「查明异象的来路」——你从看到地名的那一刻起，就知道来路了吧。', 'r3'),
        n('r3', 'test_approach', '', '', RINNE,
          '嗯。所以才接。「退魔寮乙字科，凛音——接下逢魔神社的夜巡。」这句话，我练了十年。'),
        # ---- 夜二回廊：薄荷的考据，送钟礼数 ----
        n('c1', 'corridor_act', 'enter', '', RINNE, '第二夜。回廊的灯比昨夜低了三寸——雾在往屋里走。', 'c2'),
        n('c2', 'corridor_act', '', '', HAKUYO, '钟就在头上。没有风，注连绳却是湿的……像被谁的眼泪泡过。', 'c3'),
        n('c3', 'corridor_act', '', '', HAKUYO,
          '我看过了，这口吊钟是旧式送钟——不是报时用的，是给人「送行」的。鸣三响，起灵；鸣九响，下葬。', 'c4'),
        n('c4', 'corridor_act', '', '', RINNE, '它昨夜鸣了几响？', 'c5'),
        n('c5', 'corridor_act', '', '', HAKUYO,
          '……三响。凛音，这口钟不是在报时，是在起灵。有人在这座神社里，办了一场没人看见的葬礼——办了七十年。', 'c6'),
        n('c6', 'corridor_act', '', '', RINNE, '那就让它办完。'),
        # ---- 夜三回廊深处：列队的真相 ----
        n('p1', 'night3_procession', 'enter', '', HAKUYO, '凛音，你听。纸在响——好多好多张纸，在夜里摩擦的声音。', 'p2'),
        n('p2', 'night3_procession', '', '', RINNE, '列队。纸人形抬着一口空棺，正沿着回廊往本殿去。……棺是空的，重量却在。', 'p4'),
        n('p3', 'night3_procession', 'hotspot', '700', SAYO, '（棺中传来三声轻叩。不快，不慢，像在数谁的脚步。）'),
        n('p4', 'night3_procession', '', '', HAKUYO,
          '书上写过这种列队——扎纸守灵。人走了，心里的念头没走，就会附在守灵的纸人身上，一遍一遍，继续做生前没做完的那件事。', 'p5'),
        n('p5', 'night3_procession', '', '', RINNE, '它们抬了七十年。抬棺的人不累——因为走路的早就不是人了。累的是那个下命令的。', 'p6'),
        n('p6', 'night3_procession', '', '', HAKUYO,
          '凛音，我在数——每圈列队经过，棺里都叩三下。她在催人。可她催的不是我们……是让她等的那个人。'),
        # ---- 夜四本殿黄昏：小夜开口（经薄荷转述），凛音摊牌私缘 ----
        n('m1', 'night4_mirror', 'enter', '', RINNE, '第四夜。棺停在御神木下了。树洞里那圈纹路，和棺盖上的一模一样。', 'm2'),
        n('m2', 'night4_mirror', '', '', HAKUYO, '棺里没有尸气，凛音……棺底嵌着一面镜子。镜中的人，穿着七十年前的嫁衣。', 'm3'),
        n('m3', 'night4_mirror', '', '', RINNE, '不是怨灵。四夜的异象，从头到尾，只是一个人想被好好送走。', 'm4'),
        n('m4', 'night4_mirror', '', '', HAKUYO,
          '我能听见她。很轻，像隔着很深的水……她说她叫小夜。她说等的人一直没来，问我们是不是「山下顾家派来的」。', 'm5'),
        n('m5', 'night4_mirror', '', '', RINNE,
          '（沉默了一会儿）顾家。……我祖母的娘家，就姓顾。当年订下这门亲事的，是小夜的兄长。战乱起的那年，顾家整族南迁——那盏聘礼之灯，没来得及送上山。', 'm6'),
        n('m6', 'night4_mirror', '', '', HAKUYO,
          '所以送葬的仪式差了最后一环，乡俗里她就进不了祖坟……七十年，就卡在这一盏灯上？', 'm7'),
        n('m7', 'night4_mirror', '', '', RINNE,
          '卡在「没人记得」上。灯、钟、棺、镜——她把路标一个一个点亮了。明天是逢魔刻。该有人把这盏灯，送到她面前了。'),
        # ---- 夜五本殿：代顾家叩门，棺守让路 ----
        n('hd1', 'honden_act', 'enter', '', HAKUYO, '逢魔刻。纸人形都在本殿前列队了——它们在等一个仪式开始。', 'hd2'),
        n('hd2', 'honden_act', '', '', RINNE, '也是结束。它催人三夜，叩门三声。今晚，换我叩门。', 'hd3'),
        n('hd3', 'honden_act', '', '', GUANSHOU, '来者——是送行的，还是夺棺的？', 'hd4'),
        n('hd4', 'honden_act', '', '', RINNE,
          '送行的。退魔寮乙字科，凛音——受山下顾家之托，送顾氏未过门的媳妇，小夜姑娘，上路。', 'hd5'),
        n('hd5', 'honden_act', '', '', GUANSHOU,
          '……顾。七十年了。原来山下还有人，记得这个姓。（他让开了半步）……那么，拜托了。'),
    ]
    for r in rows:
        ws.append(r)
    wb.save(f'{ROOT}/config/dialogue.xlsx')
    print('dialogue.xlsx OK')


# ---------------- commissions.json ----------------
def author_commissions():
    path = f'{ROOT}/data/commissions.json'
    d = json.load(open(path, encoding='utf-8'))
    commissions = [
        {'id': 'night_patrol_01', 'no': '第 一 夜', 'title': '逢魔神社 · 石钵的倒影',
         'client': '退魔寮 · 乙字科急递',
         'desc': '急递：逢魔神社参道，石钵倒影中现无源灯笼一盏，逢魔之刻其光渐盛。'
                 '按寮规三则——一辨来路，二退其形，三送其心。查明来路，不得妄杀。备注：该社七十年前已废。',
         'reward': '绘卷残页 ×3 · 朱印一枚',
         'stage': 'res://scenes/v3/stage.tscn', 'case': 'night_patrol', 'locked': False},
        {'id': 'night_patrol_02', 'no': '第 二 夜', 'title': '逢魔神社 · 无风之钟',
         'client': '退魔寮 · 乙字科急递',
         'desc': '续报：本殿回廊吊钟无风自鸣，每夜三响而止。随行文书考据——钟为旧式送钟，鸣三响者，起灵也。'
                 '异象有仪，非寻常妖物。留意钟上附着物。',
         'reward': '绘卷残页 ×3 · 朱印一枚',
         'stage': 'res://scenes/v3/stage_corridor.tscn', 'case': 'night_patrol', 'locked': False},
        {'id': 'night_patrol_03', 'no': '第 三 夜', 'title': '逢魔神社 · 纸人抬棺',
         'client': '退魔寮 · 乙字科急递',
         'desc': '续报：雾夜有扎纸列队，抬空棺一具夜行本殿，沿途灵气上浮。扎纸受执念驱动，'
                 '行生前未竟之事——列队即「守灵」，棺中所守何人，查明之。',
         'reward': '绘卷残页 ×3 · 朱印一枚',
         'stage': 'res://scenes/v3/stage_night3.tscn', 'case': 'night_patrol', 'locked': False},
        {'id': 'night_patrol_04', 'no': '第 四 夜', 'title': '逢魔神社 · 棺中镜',
         'client': '退魔寮 · 乙字科急递',
         'desc': '续报：空棺停于御神木下，棺底嵌铜镜一面，镜现嫁衣女子。随行文书灵听所得：'
                 '彼自名「小夜」，候山下顾家七十年。顾氏旧案，档房已无可考——只知道，灯没送到。',
         'reward': '绘卷残页 ×3 · 朱印一枚',
         'stage': 'res://scenes/v3/stage_night4.tscn', 'case': 'night_patrol', 'locked': False},
        {'id': 'night_patrol_05', 'no': '第 五 夜', 'title': '逢魔神社 · 逢魔刻',
         'client': '退魔寮 · 乙字科急递',
         'desc': '终递：逢魔之刻将至，该社百纸列队，守棺之物现身本殿。此案五日五夜，来路已明，'
                 '余下最后一则——送其心。处理方式依退魔师裁量：送行，或镇守。朱印在案。',
         'reward': '绘卷残页 ×3 · 朱印一枚',
         'stage': 'res://scenes/v3/stage_honden.tscn', 'case': 'night_patrol', 'locked': False},
    ]
    d['commissions'] = commissions
    json.dump(d, open(path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    print('commissions.json OK')


if __name__ == '__main__':
    author_battle()
    author_stage()
    author_dialogue()
    author_commissions()
    print('全部落表完成，下一步运行三个导出器')
