# -*- coding: utf-8 -*-
"""舞台/线索配置管线：config/stage_config.xlsx（人工维护）→ data/stages/*.json + data/clues/*.json。

探索侧全部配置：舞台（背景/地面剖面/玩家动画帧表/队伍）+ 画面线索（异象渲染参数/调查文案/转场战斗）。

用法：
  python tools/export_stage_config_xlsx.py            # 读取 xlsx → 校验 → 导出 JSON
  python tools/export_stage_config_xlsx.py --init     # 用 data/stages+data/clues 现有 JSON 反向建表

校验（警告当错误）：
  必填缺失 / JSON 单元格解析失败 / res:// 资源文件不存在 /
  bounds 升序两元组 / ground_profile ≥2 点且 x 单调不降 / spawn_x 在界内 /
  spirit 0-10 / 线索 x 在所属舞台界内 / 线索 battle .tscn 存在 / 线索 id 按舞台唯一
"""
import glob
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "config", "stage_config.xlsx")
STAGES_DIR = os.path.join(ROOT, "data", "stages")
CLUES_DIR = os.path.join(ROOT, "data", "clues")
CASES_DIR = os.path.join(ROOT, "data", "cases")

CLUE_FILE_COMMENT = ("画面线索（GDD §8 Phase 3）：探索中观察画面异常 → 靠近出现调查提示 → E 调查 → "
                     "解密文案 → 墨线转场进入战斗。坐标为世界坐标（与 stages/*.json 同一空间），"
                     "y 缺省时取地面剖面。")

STAGE_FIELDS = [
    ("file", "file(输出文件名)", "str"),
    ("id", "id", "str"),
    ("name", "name(场景名)", "str"),
    ("comment", "comment(备注)", "str"),
    ("bg", "bg(背景图)", "str"),
    ("bg_far", "bg_far(远景层)", "str"),
    ("bg_far_factor", "bg_far_factor(远景视差)", "float"),
    ("player_sprite", "player_sprite(玩家立绘)", "str"),
    ("player_height_px", "player_height_px(立绘高)", "int"),
    ("move_speed", "move_speed(移动速度)", "float"),
    ("spawn_x", "spawn_x(出生x)", "int"),
    ("spirit", "spirit(初始灵气)", "int"),
    ("party", "party(队伍JSON)", "json"),
    ("bounds", "bounds(活动范围JSON)", "json"),
    ("ground_profile", "ground_profile(地面剖面JSON)", "json"),
    ("player_anims", "player_anims(动画JSON)", "json"),
    ("player_anims_from", "player_anims_from(动画清单路径,优先)", "str"),
    ("next", "next(下一幕场景,可空)", "str"),
    ("comment_ground", "comment_ground(剖面备注)", "str"),
]
CLUE_FIELDS = [
    ("stage", "stage(舞台file)", "str"),
    ("idx", "idx(序)", "int"),
    ("id", "id", "str"),
    ("type", "type(异象类型)", "str"),
    ("name", "name(线索名)", "str"),
    ("x", "x(世界坐标)", "float"),
    ("y_offset", "y_offset(y偏移)", "float"),
    ("y_abs", "y_abs(锚定画中y,可空)", "float"),
    ("radius", "radius(调查半径)", "float"),
    ("anomaly", "anomaly(异象渲染JSON)", "json"),
    ("hint", "hint(靠近提示)", "str"),
    ("resolve_text", "resolve_text(解密文案)", "str"),
    ("battle", "battle(转场战斗场景)", "str"),
    ("comment", "comment(备注)", "str"),
    ("file_comment", "file_comment(文件级备注)", "str"),
]
# 案件结案（GDD §6 样章：真相→处理方式→余波→成长奖励）：nodes 为 DialogueOverlay 兼容格式
CASE_FIELDS = [
    ("file", "file(输出文件名)", "str"),
    ("id", "id", "str"),
    ("name", "name(案件名)", "str"),
    ("comment", "comment(备注)", "str"),
    ("nodes", "nodes(结案对话JSON)", "json"),
    ("rewards", "rewards(解锁奖励JSON)", "json"),
    ("stamp", "stamp(结案朱印文案)", "str"),
]
SHEETS = {"stage": STAGE_FIELDS, "clue": CLUE_FIELDS, "case": CASE_FIELDS}
REQUIRED = {
    "stage": ["file", "id", "name", "bg", "player_height_px", "spawn_x", "spirit",
              "party", "bounds", "ground_profile"],
    "clue": ["stage", "id", "name", "x", "radius", "hint", "resolve_text", "battle"],
    "case": ["file", "id", "name", "nodes", "rewards", "stamp"],
}


def fail(msg: str) -> None:
    print(f"[校验失败] {msg}")
    sys.exit(1)


def cell_to_value(raw, kind: str, sheet: str, rowno: int, key: str):
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return None, True
    where = f"{sheet} 第{rowno}行 {key}"
    try:
        if kind == "str":
            return str(raw), False
        if kind == "int":
            return int(raw), False
        if kind == "float":
            return float(raw), False
        if kind == "json":
            return json.loads(str(raw)), False
    except (ValueError, TypeError, json.JSONDecodeError) as e:
        fail(f"{where}: 解析失败（{e}）")


def read_sheet(ws, fields) -> list:
    headers = [c.value for c in ws[1]]
    col_of = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        for key, header, _kind in fields:
            if str(h).strip() == header:
                col_of[key] = i
    missing = [k for k, _h, _k in fields if k not in col_of]
    if missing:
        fail(f"sheet[{ws.title}] 缺列: {missing}")
    rows = []
    for rno, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec = {}
        empty_keys = set()
        for key, _header, kind in fields:
            v, was_empty = cell_to_value(row[col_of[key]] if col_of[key] < len(row) else None,
                                         kind, ws.title, rno, key)
            if was_empty:
                empty_keys.add(key)
            else:
                rec[key] = v
        rec["_empty"] = empty_keys
        rec["_row"] = rno
        rows.append(rec)
    return rows


def check_res_file(path: str, where: str) -> None:
    if not isinstance(path, str) or not path.startswith("res://"):
        fail(f"{where}: 资源路径必须是 res:// 开头: {path!r}")
    fs = os.path.join(ROOT, path.replace("res://", "").replace("/", os.sep))
    if not os.path.isfile(fs):
        fail(f"{where}: 资源文件不存在 {path}")


def check_anims(anims: dict, where: str) -> None:
    if not isinstance(anims, dict):
        fail(f"{where}: player_anims 必须是 JSON 对象")
    for state, frames in anims.items():
        if state == "fps":
            if not (1 <= int(frames) <= 60):
                fail(f"{where}: anims.fps 必须在 1-60")
            continue
        if not isinstance(frames, list) or not frames:
            fail(f"{where}: anims.{state} 必须是非空帧数组")
        for p in frames:
            check_res_file(p, where)


def export() -> None:
    if not os.path.isfile(XLSX):
        fail(f"找不到配置源 {XLSX}，请先运行 --init")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for name in ("stage", "clue"):
        if name not in wb.sheetnames:
            fail(f"缺少 sheet[{name}]")
    stages = read_sheet(wb["stage"], STAGE_FIELDS)
    clues = read_sheet(wb["clue"], CLUE_FIELDS)
    cases = read_sheet(wb["case"], CASE_FIELDS) if "case" in wb.sheetnames else []

    stage_bounds = {}
    for s in stages:
        rowno = s["_row"]
        for k in REQUIRED["stage"]:
            if k in s["_empty"]:
                fail(f"stage 第{rowno}行: 必填缺失 {k}")
        sid = s["id"]
        where = f"舞台[{sid}]"
        check_res_file(s["bg"], where)
        check_res_file(s.get("player_sprite", ""), where)
        if s.get("player_anims_from"):
            check_res_file(s["player_anims_from"], where)
        elif "player_anims" in s:
            check_anims(s["player_anims"], where)
        b = s["bounds"]
        if not (isinstance(b, list) and len(b) == 2 and b[0] < b[1]):
            fail(f"{where}: bounds 必须是升序两元组 [min,max]")
        gp = s["ground_profile"]
        if not (isinstance(gp, list) and len(gp) >= 2):
            fail(f"{where}: ground_profile 至少 2 个点")
        xs = [float(p[0]) for p in gp]
        if any(xs[i] > xs[i + 1] for i in range(len(xs) - 1)):
            fail(f"{where}: ground_profile 的 x 必须单调不降")
        if not (b[0] <= s["spawn_x"] <= b[1]):
            fail(f"{where}: spawn_x {s['spawn_x']} 不在 bounds {b} 内")
        if not (0 <= s["spirit"] <= 10):
            fail(f"{where}: spirit 必须在 0-10")
        if s["move_speed"] if "move_speed" in s else True:
            pass
        if "move_speed" in s and s["move_speed"] <= 0:
            fail(f"{where}: move_speed 必须为正")
        if s.get("next"):
            if not str(s["next"]).endswith(".tscn"):
                fail(f"{where}: next 必须指向 .tscn 场景")
            check_res_file(s["next"], where)
        party = s["party"]
        if not (isinstance(party, list) and party):
            fail(f"{where}: party 必须是非空数组")
        for m in party:
            check_res_file(m.get("art", ""), where)
            if int(m.get("hp_max", 0)) <= 0:
                fail(f"{where}: 队员 {m.get('name')} hp_max 必须为正")
        stage_bounds[s["file"]] = (b, where)

    out_stages = 0
    for s in stages:
        out = {k: v for k, v in s.items() if not k.startswith("_")}
        out.pop("file", None)
        if out.get("player_anims_from"):
            # 清单优先：清单内含锚点/移速/逐帧时长，旧内联字段与 move_speed 不再输出
            out.pop("player_anims", None)
            out.pop("move_speed", None)
        path = os.path.join(STAGES_DIR, s["file"] + ".json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, indent=1)
        out_stages += 1

    clue_ids = set()
    by_stage: dict = {}
    for c in clues:
        rowno = c["_row"]
        for k in REQUIRED["clue"]:
            if k in c["_empty"]:
                fail(f"clue 第{rowno}行: 必填缺失 {k}")
        sid = c["stage"]
        where = f"线索[{sid}/{c.get('id', '?')}] 第{rowno}行"
        if sid not in stage_bounds:
            fail(f"{where}: 所属舞台 {sid} 不在 stage 表中")
        b, swhere = stage_bounds[sid]
        if not (b[0] <= c["x"] <= b[1]):
            fail(f"{where}: x {c['x']} 不在舞台活动范围 {b} 内（{swhere}）")
        if c["radius"] <= 0:
            fail(f"{where}: radius 必须为正")
        if not str(c["battle"]).endswith(".tscn"):
            fail(f"{where}: battle 必须指向 .tscn 场景")
        check_res_file(c["battle"], where)
        key = (sid, c["id"])
        if key in clue_ids:
            fail(f"{where}: 线索 id 在舞台内重复 {c['id']}")
        clue_ids.add(key)
        rec = {k: v for k, v in c.items() if not k.startswith("_")}
        rec.pop("stage", None)
        rec.pop("idx", None)
        fc = rec.pop("file_comment", "")
        entry = by_stage.setdefault(sid, {"comment": fc, "clues": []})
        if fc and not entry["comment"]:
            entry["comment"] = fc
        entry["clues"].append(rec)

    out_clues = 0
    for sid, bundle in by_stage.items():
        doc = {"stage": sid,
               "comment": bundle["comment"] or CLUE_FILE_COMMENT,
               "clues": bundle["clues"]}
        path = os.path.join(CLUES_DIR, sid + ".json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(doc, f, ensure_ascii=False, indent=1)
        out_clues += 1

    print(f"[导出] {XLSX} → data/stages/({out_stages}) + data/clues/({out_clues})："
          f"{sum(len(v['clues']) for v in by_stage.values())} 条线索")

    # 案件结案：data/cases/<file>.json（D 期完整单元闭环）
    card_ids = set()
    for path in glob.glob(os.path.join(ROOT, "data", "battles", "*.json")):
        with open(path, encoding="utf-8") as f:
            for c in json.load(f).get("cards", []):
                card_ids.add(c.get("id", ""))
    art_path = os.path.join(ROOT, "data", "artifacts.json")
    artifact_ids = set()
    if os.path.isfile(art_path):
        with open(art_path, encoding="utf-8") as f:
            for a in json.load(f):
                artifact_ids.add(a.get("id", ""))
    os.makedirs(CASES_DIR, exist_ok=True)
    out_cases = 0
    for c in cases:
        rowno = c["_row"]
        for k in REQUIRED["case"]:
            if k in c["_empty"]:
                fail(f"case 第{rowno}行: 必填缺失 {k}")
        where = f"案件[{c.get('id','?')}] 第{rowno}行"
        nodes = c["nodes"]
        if not isinstance(nodes, dict) or not nodes:
            fail(f"{where}: nodes 必须是非空对象（DialogueOverlay 节点表）")
        for nid, node in nodes.items():
            if not isinstance(node, dict) or not node.get("text"):
                fail(f"{where}: 节点 {nid} 缺 text")
            nxt = node.get("next", "")
            if nxt and nxt not in nodes:
                fail(f"{where}: 节点 {nid} next 悬空引用 {nxt}")
            for ch in node.get("choices", []) or []:
                if not ch.get("text") or ch.get("next", "") not in nodes:
                    fail(f"{where}: 节点 {nid} 分支缺 text 或 next 悬空")
        rw = c["rewards"]
        for cid in rw.get("cards", []) or []:
            if cid not in card_ids:
                fail(f"{where}: 奖励卡不存在于卡库 {cid}")
        for aid in rw.get("artifacts", []) or []:
            if aid not in artifact_ids:
                fail(f"{where}: 奖励器物不存在 {aid}")
        rec = {k: v for k, v in c.items() if not k.startswith("_")}
        rec.pop("file", None)
        path = os.path.join(CASES_DIR, c["file"] + ".json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(rec, f, ensure_ascii=False, indent=1)
        out_cases += 1
    if out_cases:
        print(f"  cases/：{out_cases} 案（{', '.join(c.get('id','?') for c in cases)}）")


def build_xlsx() -> None:
    """从 data/stages/*.json + data/clues/*.json 反向生成 xlsx（无损）"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3B3B58")

    def put(ws, fields, rows):
        for ci, (_k, header, _t) in enumerate(fields, 1):
            cell = ws.cell(1, ci, header)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(ci)].width = max(12, min(46, len(header) + 6))
        for ri, rec in enumerate(rows, 2):
            for ci, (key, _h, kind) in enumerate(fields, 1):
                if key not in rec:
                    continue
                v = rec[key]
                if kind == "json":
                    v = json.dumps(v, ensure_ascii=False)
                cell = ws.cell(ri, ci, v)
                if isinstance(v, str) and len(v) > 24:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"

    stage_rows, clue_rows = [], []
    for path in sorted(glob.glob(os.path.join(STAGES_DIR, "*.json"))):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        fname = os.path.splitext(os.path.basename(path))[0]
        row = {"file": fname}
        row.update(data)
        stage_rows.append(row)
    for path in sorted(glob.glob(os.path.join(CLUES_DIR, "*.json"))):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        for i, c in enumerate(data.get("clues", [])):
            row = {"stage": data.get("stage", ""), "idx": i,
                   "file_comment": data.get("comment", "")}
            row.update(c)
            clue_rows.append(row)
    case_rows = []
    for path in sorted(glob.glob(os.path.join(CASES_DIR, "*.json"))):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        row = {"file": os.path.splitext(os.path.basename(path))[0]}
        row.update(data)
        case_rows.append(row)
    for name, fields, rows in (("stage", STAGE_FIELDS, stage_rows),
                               ("clue", CLUE_FIELDS, clue_rows),
                               ("case", CASE_FIELDS, case_rows)):
        ws = wb.create_sheet(name)
        put(ws, fields, rows)
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    print(f"[init] 已生成 {XLSX}：{len(stage_rows)} 舞台 / {len(clue_rows)} 线索 / "
          f"{len(case_rows)} 案件")


if __name__ == "__main__":
    if "--init" in sys.argv:
        build_xlsx()
    else:
        export()
