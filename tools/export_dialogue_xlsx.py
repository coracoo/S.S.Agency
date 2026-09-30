# -*- coding: utf-8 -*-
"""对话台词配置管线：config/dialogue.xlsx（人工维护的工程化配置源）→ data/dialogues.json（游戏运行时读取）。

用法：
  python tools/export_dialogue_xlsx.py            # 读取 xlsx → 校验 → 导出 JSON
  python tools/export_dialogue_xlsx.py --init     # 用内置剧情内容重建 xlsx（首次/重置用）

xlsx 结构（两个 sheet）：
  nodes   : node_id | stage | trigger | trigger_arg | speaker | display_name | color
            | portrait | side | text | next
            trigger = enter（进场景自动播）/ hotspot（trigger_arg=世界 x 坐标，走到触发）
            next 为空 = 链结束；节点配了 choices 时 next 忽略
  choices : choice_id | node_id | text | next

校验（警告当错误，退出码非 0）：
  node_id 重复 / next 与 choices.next 悬空引用 / trigger 非法 / hotspot 坐标非数字
  / stage 为空 / portrait 文件不存在
"""
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "config", "dialogue.xlsx")
OUT = os.path.join(ROOT, "data", "dialogues.json")

# 内置剧情内容：第一夜·参道（test_approach）。改台词请改 xlsx，不要改这里后忘记同步；
# --init 仅用于首次生成模板/整体重置。
NODES = [
    # node_id, stage, trigger, trigger_arg, speaker, display_name, color, portrait, side, text, next
    ["a1", "test_approach", "enter", "", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "就是这条参道了。逢魔时刻还差一点——山里的雾，已经开始发烫。", "a2"],
    ["a2", "test_approach", "", "", "hakuyo", "薄荷", "fuji_500",
     "res://assets/chars/hakuyo_idle.png", "right",
     "委托人说的「石钵倒影」，真的会在日落时分亮起来吗？我现在可是什么都看不见哦。", "a3"],
    ["a3", "test_approach", "", "", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "看不见才好。看得一清二楚的东西，事务所收三倍价钱。", "a4"],
    ["a4", "test_approach", "", "", "hakuyo", "薄荷", "fuji_500",
     "res://assets/chars/hakuyo_idle.png", "right",
     "诶——那上次那口井里的，你怎么只收了一倍？", "a5"],
    ["a5", "test_approach", "", "", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "因为那口井欠我一次。……到了。台阶上面，就是村子里没人敢提的那块地。", ""],
    ["h1", "test_approach", "hotspot", "850", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "停。薄荷，看台阶边上那尊石钵——水面上有光。", "h2"],
    ["h2", "test_approach", "", "", "hakuyo", "薄荷", "fuji_500",
     "res://assets/chars/hakuyo_idle.png", "right",
     "冷冷的……蓝紫色的。可是凛音，那倒影里——没有我们。", "h3"],
    ["h3", "test_approach", "", "", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "石钵照的不是人，是「缝隙」。水面亮起来的时候，逢魔的眼睛也睁开了。", ""],
    ["h4a", "test_approach", "", "", "hakuyo", "薄荷", "fuji_500",
     "res://assets/chars/hakuyo_idle.png", "right",
     "（压低声音）光是从水底浮上来的……像灯笼火，可是水下根本没有灯。", "h5"],
    ["h4b", "test_approach", "", "", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "别碰水面。……记住了：它现在只是「看」着我们，还不想把人拉下去。", "h5"],
    ["h5", "test_approach", "", "", "rinne", "凛音", "vermilion_500",
     "res://assets/chars/rinne_idle.png", "left",
     "靠近石钵，按 E 确认那道倒影。看完就走——天黑之前，我们还有一场架要打。", ""],
]

CHOICES = [
    # choice_id, node_id, text, next
    ["c1", "h3", "先退后一步，看清光的来路", "h4a"],
    ["c2", "h3", "压低脚步，直接靠近", "h4b"],
]

NODE_HEADERS = ["node_id", "stage", "trigger", "trigger_arg", "speaker",
                "display_name", "color", "portrait", "side", "text", "next"]
CHOICE_HEADERS = ["choice_id", "node_id", "text", "next"]


def build_xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "nodes"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3B3B58")
    for c, h in enumerate(NODE_HEADERS, 1):
        cell = ws.cell(1, c, h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(NODES, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v).alignment = Alignment(vertical="top", wrap_text=(c == 10))
    widths = [10, 16, 10, 12, 10, 12, 14, 34, 8, 60, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("choices")
    for c, h in enumerate(CHOICE_HEADERS, 1):
        cell = ws2.cell(1, c, h)
        cell.font = head_font
        cell.fill = head_fill
    for r, row in enumerate(CHOICES, 2):
        for c, v in enumerate(row, 1):
            ws2.cell(r, c, v).alignment = Alignment(vertical="top", wrap_text=(c == 3))
    for i, w in enumerate([10, 10, 40, 8], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    print(f"[init] 已生成 {XLSX}（{len(NODES)} 节点 / {len(CHOICES)} 选项）")


def fail(msg: str) -> None:
    print(f"[校验失败] {msg}")
    sys.exit(1)


def check_portrait(path: str) -> None:
    if not path.startswith("res://"):
        fail(f"portrait 必须是 res:// 路径: {path}")
    fs = os.path.join(ROOT, path.replace("res://", "").replace("/", os.sep))
    if not os.path.isfile(fs):
        fail(f"portrait 文件不存在: {path}")


def export() -> None:
    if not os.path.isfile(XLSX):
        fail(f"找不到配置源 {XLSX}，请先运行 --init")
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # ---- nodes ----
    ws = wb["nodes"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    nodes: dict = {}
    for row in rows:
        if row[0] is None or str(row[0]).strip() == "":
            continue
        rec = dict(zip(NODE_HEADERS, ["" if v is None else str(v).strip() for v in row]))
        nid = rec["node_id"]
        if nid in nodes:
            fail(f"node_id 重复: {nid}")
        if rec["stage"] == "":
            fail(f"{nid}: stage 为空")
        if rec["trigger"] not in ("", "enter", "hotspot"):
            fail(f"{nid}: 非法 trigger '{rec['trigger']}'")
        if rec["trigger"] == "hotspot":
            try:
                float(rec["trigger_arg"])
            except ValueError:
                fail(f"{nid}: hotspot 坐标非数字 '{rec['trigger_arg']}'")
        if rec["text"] == "":
            fail(f"{nid}: text 为空")
        if rec["side"] not in ("left", "right"):
            fail(f"{nid}: side 必须是 left/right")
        if rec["portrait"]:
            check_portrait(rec["portrait"])
        nodes[nid] = {
            "speaker": rec["speaker"],
            "name": rec["display_name"],
            "color": rec["color"],
            "portrait": rec["portrait"],
            "side": rec["side"],
            "text": rec["text"],
            "next": rec["next"],
            "trigger": rec["trigger"],
            "trigger_arg": rec["trigger_arg"],
            "stage": rec["stage"],
            "choices": [],
        }

    # ---- choices ----
    if "choices" in wb.sheetnames:
        for row in wb["choices"].iter_rows(min_row=2, values_only=True):
            if row[0] is None or str(row[0]).strip() == "":
                continue
            cid, node_id, text, nxt = ["" if v is None else str(v).strip() for v in row[:4]]
            if node_id not in nodes:
                fail(f"选项 {cid}: 挂在不存在的节点 {node_id}")
            nodes[node_id]["choices"].append({"text": text, "next": nxt})

    # ---- 引用完整性 ----
    for nid, n in nodes.items():
        refs = [c["next"] for c in n["choices"] if c["next"]]
        if n["choices"] and not refs:
            fail(f"{nid}: 有选项但所有 next 为空")
        if n["next"]:
            refs.append(n["next"])
        for r in refs:
            if r not in nodes:
                fail(f"{nid}: 悬空引用 next='{r}'")

    # ---- 按 stage 聚合（触发索引 + 节点表）----
    stages: dict = {}
    for nid, n in nodes.items():
        sid = n["stage"]
        st = stages.setdefault(sid, {"enter": "", "hotspots": [], "nodes": {}})
        st["nodes"][nid] = {k: n[k] for k in
                            ("speaker", "name", "color", "portrait", "side", "text", "next", "choices")}
        if n["trigger"] == "enter":
            if st["enter"]:
                fail(f"stage {sid}: enter 触发器只能有一个（已有 {st['enter']}，又遇 {nid}）")
            st["enter"] = nid
        elif n["trigger"] == "hotspot":
            st["hotspots"].append({"x": float(n["trigger_arg"]), "node": nid})
    for st in stages.values():
        st["hotspots"].sort(key=lambda h: h["x"])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump({"stages": stages}, f, ensure_ascii=False, indent=1)
    n_hot = sum(len(s["hotspots"]) for s in stages.values())
    print(f"[导出] {OUT}：{len(stages)} 舞台 / {len(nodes)} 节点 / "
          f"{sum(len(n['choices']) for s in stages.values() for n in s['nodes'].values())} 选项 / {n_hot} 热点")


if __name__ == "__main__":
    if "--init" in sys.argv:
        build_xlsx()
    else:
        export()
